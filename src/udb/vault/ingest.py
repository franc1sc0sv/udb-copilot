"""Turn source files into citable text (`derived/text/<ID>.pages.jsonl`) and index it.

Every unit Claude can cite is a "page" record:

- PDF page or slide: its text; a scanned page gets OCR text (`"ocr": true`).
- HTML: one record per section (h1-h3), with the section title, so a citation
  points at "Inverso multiplicativo" instead of the whole site.
- Figure of an HTML page: the OCR text of the image (`"kind": "figure"`), added
  by `fetch_images` once the image is downloaded.

`derived/text/<ID>.meta.json` keeps what could not become text yet (images not
downloaded, embedded videos and documents, pages that need OCR), so the source
report can say exactly what is missing instead of "I can't read it".
"""
import json
import re
from pathlib import Path
from urllib.parse import urljoin

from ..core.paths import derived_dir
from . import index, ocr, registry

EXTRACTOR = "udb-ingest/3"
TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".rst"}
AUDIO_SUFFIXES = {".m4a", ".mp3", ".wav", ".mp4", ".mov", ".mkv", ".webm"}
UNSUPPORTED = {".rar", ".zip", ".7z", ".exe", ".pkt", ".slx"}
EMPTY_PAGE = 30          # characters; less than this is "no text"
SECTION_TAGS = ["h1", "h2", "h3"]
MIN_SECTION = 80         # tiny sections merge into the previous one
SECTION_MARK = "\u0000SECTION\u0000"
CAPTION_RE = re.compile(r"(?i)^(figura|gr[aá]fic[oa]|imagen|tabla|figure)\s*\d+")


def meta_file(code: str, source_id: str) -> Path:
    return derived_dir(code) / "text" / f"{source_id}.meta.json"


def images_dir(code: str, source_id: str) -> Path:
    return derived_dir(code) / "images" / source_id


def load_meta(code: str, source_id: str) -> dict:
    path = meta_file(code, source_id)
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}


def save_meta(code: str, source_id: str, meta: dict) -> None:
    path = meta_file(code, source_id)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(meta, ensure_ascii=False, indent=1), encoding="utf-8")


# ------------------------------------------------------------------ extractors

def extract_pdf(path: Path, use_ocr: bool = True) -> tuple[list[dict], dict]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages, needs_ocr = [], []
    can_ocr = use_ocr and ocr.available()
    for i, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        record = {"page": i, "text": text}
        if len(text) < EMPTY_PAGE:
            images = list(page.images)
            if images and can_ocr:
                record["text"] = "\n".join(t for t in (ocr.read(img.data) for img in images) if t)
                record["ocr"] = True
            elif images:
                needs_ocr.append(i)
        pages.append(record)
    return pages, {"needs_ocr": needs_ocr}


def extract_pptx(path: Path, use_ocr: bool = True) -> tuple[list[dict], dict]:
    from pptx import Presentation

    pages = []
    for i, slide in enumerate(Presentation(str(path)).slides, 1):
        parts = [sh.text_frame.text for sh in slide.shapes
                 if sh.has_text_frame and sh.text_frame.text.strip()]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append(slide.notes_slide.notes_text_frame.text)
        pages.append({"page": i, "text": "\n".join(parts).strip(), "kind": "slide"})
    return pages, {}


def extract_docx(path: Path, use_ocr: bool = True) -> tuple[list[dict], dict]:
    import docx

    blocks: list[tuple[str, list[str]]] = [("", [])]
    for para in docx.Document(str(path)).paragraphs:
        if para.style.name.startswith("Heading") and blocks[-1][1]:
            blocks.append((para.text.strip(), []))
        if para.text.strip():
            blocks[-1][1].append(para.text)
    return [{"page": i, "text": "\n".join(lines), "kind": "section", "section": title}
            for i, (title, lines) in enumerate((b for b in blocks if b[1]), 1)], {}


def _caption(img) -> str:
    for node in [img, *img.parents]:
        for sib in (node.find_next_sibling(), node.find_previous_sibling()):
            if sib is not None and hasattr(sib, "get_text"):
                txt = sib.get_text(" ", strip=True)
                if CAPTION_RE.match(txt):
                    return txt[:200]
        alt = node.attrs.get("alt") if hasattr(node, "attrs") else None
        if alt and alt.strip():
            return alt.strip()[:200]
    return ""


def extract_html(path: Path, use_ocr: bool = True) -> tuple[list[dict], dict]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "noscript"]):
        tag.decompose()

    images, embeds = [], []
    for n, img in enumerate(soup.find_all("img"), 1):
        caption = _caption(img)
        images.append({"n": n, "src": img.get("src") or "", "caption": caption,
                       "status": "pending"})
        img.replace_with(f" [Figura {n}{': ' + caption if caption else ''}] ")
    for frame in soup.find_all(["iframe", "embed", "object"]):
        src = frame.get("src") or frame.get("data") or ""
        title = frame.get("title") or ""
        embeds.append({"src": src, "title": title})
        frame.replace_with(f" [Contenido incrustado{': ' + title if title else ''} — {src}] ")

    for heading in soup.find_all(SECTION_TAGS):
        heading.insert_before(SECTION_MARK)
    raw = soup.get_text("\n", strip=True)
    parts = [p.strip() for p in raw.split(SECTION_MARK)]

    sections: list[dict] = []
    for part in filter(None, parts):
        # A heading with almost nothing under it joins the section before it, so a
        # citation never points at a two-word "page"; the first section always stays.
        if sections and len(part) < MIN_SECTION:
            sections[-1]["text"] += "\n" + part
            continue
        sections.append({"text": part, "kind": "section", "section": part.split("\n", 1)[0][:120]})
    pages = [{"page": i, **s} for i, s in enumerate(sections, 1)] or [{"page": 1, "text": ""}]
    return pages, {"images": images, "embeds": embeds}


def extract_text(path: Path, use_ocr: bool = True) -> tuple[list[dict], dict]:
    parts = [p.strip() for p in path.read_text(encoding="utf-8", errors="replace").split("\n\n")
             if p.strip()]
    return [{"page": i, "text": p} for i, p in enumerate(parts, 1)], {}


EXTRACTORS = {".pdf": extract_pdf, ".pptx": extract_pptx, ".docx": extract_docx,
              ".html": extract_html, ".htm": extract_html}


def extract(path: Path, use_ocr: bool = True) -> tuple[list[dict], dict]:
    suffix = path.suffix.lower()
    if suffix in EXTRACTORS:
        return EXTRACTORS[suffix](path, use_ocr)
    if suffix in TEXT_SUFFIXES:
        return extract_text(path, use_ocr)
    raise ValueError(f"unsupported file type: {suffix}")


# ------------------------------------------------------------------ writing and indexing

def write_pages(code: str, source_id: str, pages: list[dict]) -> Path:
    out = index.pages_file(code, source_id)
    out.parent.mkdir(parents=True, exist_ok=True)
    cursor, lines = 0, []
    for page in pages:
        text = page.get("text", "")
        record = {**page, "char_start": cursor, "char_end": cursor + len(text)}
        lines.append(json.dumps(record, ensure_ascii=False))
        cursor += len(text) + 1
    out.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return out


def reindex_source(code: str, source_id: str) -> int:
    con = index.connect()
    try:
        index.upsert_source(con, registry.get(code, source_id) or {})
        index.clear_source(con, source_id)
        count = index.add_chunks(con, source_id, code.upper(),
                                 index.chunks_for(code, source_id, index.boilerplate(code)))
        con.commit()
    finally:
        con.close()
    return count


def index_source(code: str, source_id: str, use_ocr: bool = True) -> int:
    path = registry.source_file(code, source_id)
    suffix = path.suffix.lower()
    if suffix in AUDIO_SUFFIXES:
        if not index.transcript_file(code, source_id).exists():
            raise RuntimeError(f"{source_id} is audio and has no transcript. "
                               "Run: udb recording transcribe <id> (see udb recording board)")
    else:
        pages, meta = extract(path, use_ocr)
        old = load_meta(code, source_id)
        # Keep figures already downloaded; their pages are appended again below.
        done = {i["n"]: i for i in old.get("images", []) if i.get("status") == "done"}
        for img in meta.get("images", []):
            if img["n"] in done and done[img["n"]].get("src") == img["src"]:
                img.update(done[img["n"]])
        pages += figure_pages(meta, start=len(pages) + 1)
        write_pages(code, source_id, pages)
        save_meta(code, source_id, {"extractor": EXTRACTOR, **meta})
    return reindex_source(code, source_id)


def figure_pages(meta: dict, start: int) -> list[dict]:
    out = []
    for img in meta.get("images", []):
        if img.get("status") == "done" and img.get("text"):
            label = f"Figura {img['n']}" + (f". {img['caption']}" if img.get("caption") else "")
            out.append({"page": start + len(out), "kind": "figure", "figure": img["n"],
                        "section": label, "text": f"{label}\n{img['text']}"})
    return out


def fetch_images(code: str, source_id: str, fetch=None, limit: int | None = None) -> dict:
    """Download the images of an HTML source, read their text, make them citable."""
    from .figures import http_fetch

    fetch = fetch or http_fetch
    rec = registry.get(code, source_id) or {}
    meta = load_meta(code, source_id)
    if not meta.get("images"):
        return {"source_id": source_id, "done": 0, "failed": 0, "pending": 0}
    can_ocr = ocr.available()
    folder = images_dir(code, source_id)
    folder.mkdir(parents=True, exist_ok=True)
    done = failed = 0
    for img in meta["images"][:limit]:
        if img.get("status") == "done":
            continue
        src = img.get("src", "")
        url = urljoin(rec.get("url") or "", src) if not re.match(r"https?://", src) else src
        if not re.match(r"https?://", url):
            img.update(status="failed", why="relative link and the source has no URL")
            failed += 1
            continue
        try:
            data, ext = fetch(url)
        except Exception as exc:  # noqa: BLE001 - one broken image is reported, not fatal
            img.update(status="failed", why=f"{type(exc).__name__}: {exc}"[:160])
            failed += 1
            continue
        name = f"img-{img['n']:03d}.{ext}"
        (folder / name).write_bytes(data)
        img.update(file=name, status="done" if can_ocr else "downloaded",
                   text=ocr.read(data) if can_ocr and ext != "svg" else "")
        done += 1
    save_meta(code, source_id, meta)
    if done:
        index_source(code, source_id)
    return {"source_id": source_id, "done": done, "failed": failed,
            "pending": sum(1 for i in meta["images"] if i.get("status") == "pending")}


def ingest_file(code: str, path: str | Path, kind: str = "course", title: str | None = None,
                class_number: int | None = None, url: str | None = None) -> dict:
    rec = registry.register(code, Path(path), kind, title=title, class_number=class_number,
                            url=url)
    return {"source": rec, "chunks": index_source(code, rec["id"])}
